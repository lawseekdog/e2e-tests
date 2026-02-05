from __future__ import annotations

from pathlib import Path

import pytest

from tests.lawyer_workbench._support.docx import (
    assert_docx_contains,
    assert_docx_has_no_template_placeholders,
    extract_docx_text,
)
from tests.lawyer_workbench._support.flow_runner import WorkbenchFlow
from tests.lawyer_workbench._support.sse import assert_visible_response
from tests.lawyer_workbench._support.utils import unwrap_api_response


def _case_facts() -> str:
    return (
        "原告：张三E2E_MULTI。\n"
        "被告：李四E2E_MULTI。\n"
        "案由：民间借贷纠纷。\n"
        "事实：2023-01-01，被告向原告借款人民币100000元，约定2023-12-31前归还；原告已通过银行转账交付。\n"
        "到期后被告未还，原告多次催收无果。\n"
        "证据：借条、转账记录、聊天记录。\n"
        "诉求：返还本金100000元，并按年利率6%支付逾期利息，承担诉讼费。"
    )


def _assert_strategy_report_tables_filled(docx_bytes: bytes) -> None:
    """Strategy report must fill stage plan + pricing tables (no blank rows)."""
    from io import BytesIO

    from docx import Document

    doc = Document(BytesIO(docx_bytes))

    found_plan = False
    found_pricing = False

    for t in doc.tables or []:
        if not t.rows:
            continue
        header = [c.text.strip() for c in t.rows[0].cells]

        # Stage plan table (standard base: 4 phases).
        if any("阶段" in x for x in header) and any("时间" in x for x in header) and any("核心任务" in x for x in header):
            found_plan = True
            for ridx, row in enumerate(t.rows[1:5], start=1):
                cells = [c.text.strip() for c in row.cells]
                if len(cells) < 3 or not cells[0] or not cells[1] or not cells[2]:
                    raise AssertionError(f"stage plan table has blank cell at row {ridx}: {cells}")
            continue

        # Pricing table (standard base: 2 fee rows).
        if any("收费类型" in x for x in header) and any("收费金额" in x for x in header) and any("支付方式" in x for x in header):
            found_pricing = True
            for ridx, row in enumerate(t.rows[1:3], start=1):
                cells = [c.text.strip() for c in row.cells]
                if len(cells) < 2 or not cells[0] or not cells[1]:
                    raise AssertionError(f"pricing table has blank type/amount at row {ridx}: {cells}")
            continue

    assert found_plan, "stage plan table not found in strategy report docx"
    assert found_pricing, "pricing table not found in strategy report docx"


@pytest.mark.e2e
@pytest.mark.slow
async def test_civil_prosecution_private_lending_generates_multiple_documents(lawyer_client):
    evidence_dir = Path(__file__).resolve().parent / "evidence"
    paths = [
        evidence_dir / "iou.txt",
        evidence_dir / "sample_transfer_record.txt",
        evidence_dir / "sample_chat_record.txt",
    ]

    uploaded_file_ids: list[str] = []
    for p in paths:
        up = await lawyer_client.upload_file(str(p), purpose="consultation")
        fid = str(((up.get("data") or {}) if isinstance(up, dict) else {}).get("id") or "").strip()
        assert fid, f"upload failed: {up}"
        uploaded_file_ids.append(fid)

    sess = await lawyer_client.create_session(service_type_id="civil_prosecution", client_role="plaintiff")
    session_id = str(((sess.get("data") or {}) if isinstance(sess, dict) else {}).get("id") or "").strip()
    assert session_id, sess

    # Workbench case_analysis defaults (plaintiff): complaint + evidence list + strategy report.
    selected_docs = ["civil_complaint", "evidence_list_doc", "litigation_strategy_report"]

    flow = WorkbenchFlow(
        client=lawyer_client,
        session_id=session_id,
        uploaded_file_ids=uploaded_file_ids,
        overrides={},
    )

    first_sse = await flow.nudge(_case_facts(), attachments=uploaded_file_ids, max_loops=80)
    assert_visible_response(first_sse)

    async def _all_deliverables_ready(f: WorkbenchFlow) -> bool:
        await f.refresh()
        if not f.matter_id:
            return False
        for key in selected_docs:
            resp = await f.client.list_deliverables(f.matter_id, output_key=key)
            data = unwrap_api_response(resp)
            items = (data.get("deliverables") if isinstance(data, dict) else None) or []
            if not items:
                return False
            item0 = items[0] if isinstance(items[0], dict) else {}
            if not str(item0.get("file_id") or "").strip():
                return False
        return True

    await flow.run_until(_all_deliverables_ready, max_steps=120, description="multiple docx deliverables")
    assert flow.matter_id, "session did not bind to matter_id"

    # Download each docx and run basic quality checks (no placeholder leaks + parties show up).
    for out_key in selected_docs:
        dels_resp = await lawyer_client.list_deliverables(flow.matter_id, output_key=out_key)
        dels = unwrap_api_response(dels_resp)
        items = (dels.get("deliverables") if isinstance(dels, dict) else None) or []
        assert items, dels_resp
        d0 = items[0] if isinstance(items[0], dict) else {}
        file_id = str(d0.get("file_id") or "").strip()
        assert file_id, d0

        docx_bytes = await lawyer_client.download_file_bytes(file_id)
        text = extract_docx_text(docx_bytes)
        assert_docx_has_no_template_placeholders(text)
        assert_docx_contains(text, must_include=["张三E2E_MULTI", "李四E2E_MULTI"])

        # Spot-check per doc type (cheap sanity).
        if out_key == "civil_complaint":
            assert any(x in text for x in ["100000", "100,000", "10万元", "10万"]), text[:2000]
        elif out_key == "litigation_strategy_report":
            assert any(x in text for x in ["诉讼策略报告", "策略"]), text[:2000]
            # Regression guard: private lending strategy reports should not drift into admin/public-interest wording.
            assert "行政公益诉讼" not in text, text[:3000]
            _assert_strategy_report_tables_filled(docx_bytes)
        elif out_key == "evidence_list_doc":
            assert any(x in text for x in ["证据目录", "证据"]), text[:2000]
        # No extra document types in the default workbench recommendation set.
